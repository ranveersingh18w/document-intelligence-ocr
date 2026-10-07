"""
DeepSeek-OCR-2 - Python UI (Tkinter - No HTML)
Pure Python desktop application with GUI
"""

# Standard library imports
import os           # For file path operations
import json         # For JSON formatting and output
import re           # For regular expressions to parse OCR output
import argparse
import time
from pathlib import Path
import threading    # For background model loading without freezing UI

# GUI library (comes with Python)
import tkinter as tk                                    # Core GUI framework
from tkinter import ttk                                 # Themed widgets (modern look)
from tkinter import filedialog                         # File open/save dialogs
from tkinter import scrolledtext                       # Text area with scrollbar
from tkinter import messagebox                         # Alert/error popups

# Image processing library
from PIL import Image, ImageTk                         # PIL/Pillow - Load and display images in GUI

# PDF processing library
import fitz                                            # PyMuPDF - Convert PDF pages to images

# Excel export library
from openpyxl import Workbook                          # openpyxl - Write Excel .xlsx files
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side  # Excel styling

# Deep learning libraries
import torch                                           # PyTorch - GPU acceleration and model loading
from transformers import AutoModel, AutoTokenizer      # HuggingFace - Load pre-trained OCR model

PROJECT_DIR = Path(__file__).resolve().parent
MODEL_PATH = str(PROJECT_DIR / "models" / "DeepSeek-OCR-2")

class OCRApp:
    def __init__(self, root, output_dir=None):
        self.root = root
        self.output_dir = Path(output_dir).resolve() if output_dir else None
        self.root.title("Document Intelligence | Anveer Singh")
        self.root.geometry("1320x850")
        self.root.minsize(1100, 720)
        self.root.configure(bg="#eef2f6")
        
        self.model = None
        self.tokenizer = None
        self.current_image = None
        self.image_path = None
        self.pdf_pages = []          # List of image paths from PDF pages
        self.current_page = 0        # Current page index for display
        
        self.setup_ui()
        self.load_model_async()
    
    def setup_ui(self):
        """Create the user interface"""
        style = ttk.Style(self.root)
        style.theme_use("clam")
        style.configure("TFrame", background="#eef2f6")
        style.configure("TLabelframe", background="#eef2f6")
        style.configure("TLabelframe.Label", background="#eef2f6", foreground="#17324d", font=("Segoe UI", 11, "bold"))
        style.configure("TButton", font=("Segoe UI", 10), padding=(10, 7))
        style.configure("TNotebook.Tab", font=("Segoe UI", 10), padding=(16, 8))
        style.configure("TLabel", background="#eef2f6", font=("Segoe UI", 10))
        header = tk.Frame(self.root, bg="#102b46", padx=24, pady=16)
        header.pack(fill=tk.X)
        tk.Label(header, text="DOCUMENT INTELLIGENCE", font=("Segoe UI", 23, "bold"), bg="#102b46", fg="white").pack(anchor="w")
        tk.Label(header, text="PDF to structured data  |  Powered by DeepSeek-OCR-2", font=("Segoe UI", 11), bg="#102b46", fg="#a9dbe6").pack(anchor="w", pady=(4, 0))
        tk.Label(header, text="ANVEER SINGH  /  AI / ML ENGINEER INTERN", font=("Segoe UI", 10, "bold"), bg="#102b46", fg="#e2eaf3").pack(anchor="e")

        # Main container
        main_frame = ttk.Frame(self.root, padding="10")
        main_frame.pack(fill=tk.BOTH, expand=True)
        
        # Left panel - Document preview
        left_frame = ttk.LabelFrame(main_frame, text="Document Preview", padding="10")
        left_frame.grid(row=0, column=0, sticky="nsew", padx=5)
        
        # Image display
        self.image_label = tk.Label(left_frame, text="No document loaded", bg="#e2e8f0", width=40, height=20)
        self.image_label.pack(fill=tk.BOTH, expand=True, pady=5)
        
        # Page navigation frame (for multi-page PDFs)
        nav_frame = ttk.Frame(left_frame)
        nav_frame.pack(fill=tk.X, pady=2)
        
        self.prev_btn = ttk.Button(nav_frame, text="◀ Prev", command=self.prev_page, state=tk.DISABLED)
        self.prev_btn.pack(side=tk.LEFT, padx=5)
        
        self.page_label_var = tk.StringVar(value="")
        self.page_label = ttk.Label(nav_frame, textvariable=self.page_label_var)
        self.page_label.pack(side=tk.LEFT, expand=True)
        
        self.next_btn = ttk.Button(nav_frame, text="Next ▶", command=self.next_page, state=tk.DISABLED)
        self.next_btn.pack(side=tk.RIGHT, padx=5)
        
        # Buttons
        btn_frame = ttk.Frame(left_frame)
        btn_frame.pack(fill=tk.X, pady=5)
        
        self.upload_btn = ttk.Button(btn_frame, text="📁 Upload PDF", command=self.upload_pdf)
        self.upload_btn.pack(side=tk.LEFT, padx=5)
        
        self.process_btn = ttk.Button(btn_frame, text="🚀 Process All Pages", command=self.process_image, state=tk.DISABLED)
        self.process_btn.pack(side=tk.LEFT, padx=5)
        
        self.clear_btn = ttk.Button(btn_frame, text="🗑️ Clear", command=self.clear_all)
        self.clear_btn.pack(side=tk.LEFT, padx=5)
        
        # Right panel - Tabbed Output
        right_frame = ttk.LabelFrame(main_frame, text="Output (3 Formats)", padding="10")
        right_frame.grid(row=0, column=1, sticky="nsew", padx=5)
        
        # Create Notebook (Tabbed interface)
        self.notebook = ttk.Notebook(right_frame)
        self.notebook.pack(fill=tk.BOTH, expand=True, pady=5)
        
        # Tab 1: RAW OUTPUT
        raw_frame = ttk.Frame(self.notebook)
        self.notebook.add(raw_frame, text="1️⃣ Raw Output")
        
        self.raw_output = scrolledtext.ScrolledText(
            raw_frame,
            wrap=tk.WORD,
            font=("Consolas", 10),
            bg="#f5f5f5"
        )
        self.raw_output.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        
        raw_btn_frame = ttk.Frame(raw_frame)
        raw_btn_frame.pack(fill=tk.X, padx=5, pady=5)
        ttk.Button(raw_btn_frame, text="📋 Copy Raw", command=lambda: self.copy_tab("raw")).pack(side=tk.LEFT, padx=2)
        ttk.Button(raw_btn_frame, text="💾 Save Raw", command=lambda: self.save_tab("raw")).pack(side=tk.LEFT, padx=2)
        
        # Tab 2: CLEAN MARKDOWN/HTML
        clean_frame = ttk.Frame(self.notebook)
        self.notebook.add(clean_frame, text="2️⃣ Clean Markdown")
        
        self.clean_output = scrolledtext.ScrolledText(
            clean_frame,
            wrap=tk.WORD,
            font=("Segoe UI", 11),
            bg="#ffffff"
        )
        self.clean_output.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        
        clean_btn_frame = ttk.Frame(clean_frame)
        clean_btn_frame.pack(fill=tk.X, padx=5, pady=5)
        ttk.Button(clean_btn_frame, text="📋 Copy Markdown", command=lambda: self.copy_tab("clean")).pack(side=tk.LEFT, padx=2)
        ttk.Button(clean_btn_frame, text="💾 Save Markdown", command=lambda: self.save_tab("clean")).pack(side=tk.LEFT, padx=2)
        
        # Tab 3: JSON OUTPUT
        json_frame = ttk.Frame(self.notebook)
        self.notebook.add(json_frame, text="3️⃣ JSON Data")
        
        self.json_output = scrolledtext.ScrolledText(
            json_frame,
            wrap=tk.WORD,
            font=("Consolas", 10),
            bg="#ffffff"
        )
        self.json_output.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        
        json_btn_frame = ttk.Frame(json_frame)
        json_btn_frame.pack(fill=tk.X, padx=5, pady=5)
        ttk.Button(json_btn_frame, text="📋 Copy JSON", command=lambda: self.copy_tab("json")).pack(side=tk.LEFT, padx=2)
        ttk.Button(json_btn_frame, text="💾 Save JSON", command=lambda: self.save_tab("json")).pack(side=tk.LEFT, padx=2)
        ttk.Button(json_btn_frame, text="📊 Export Excel", command=self.export_excel).pack(side=tk.LEFT, padx=2)
        
        # Status bar
        self.status_var = tk.StringVar(value="Ready")
        status_bar = ttk.Label(self.root, textvariable=self.status_var, relief=tk.SUNKEN, anchor=tk.W)
        status_bar.pack(side=tk.BOTTOM, fill=tk.X)
        
        # Configure grid weights
        main_frame.columnconfigure(0, weight=1)
        main_frame.columnconfigure(1, weight=2)
        main_frame.rowconfigure(0, weight=1)
    
    def load_model_async(self):
        """Load model in background"""
        def load():
            self.status_var.set("Loading model...")
            try:
                if not torch.cuda.is_available():
                    self.status_var.set("CUDA unavailable - install a CUDA-enabled PyTorch build")
                    messagebox.showerror("Error", "CUDA not available! Please check your PyTorch installation.")
                    return
                
                tokenizer = AutoTokenizer.from_pretrained(MODEL_PATH, trust_remote_code=True)
                dtype = torch.bfloat16 if torch.cuda.is_bf16_supported() else torch.float16
                
                try:
                    model = AutoModel.from_pretrained(
                        MODEL_PATH,
                        attn_implementation='flash_attention_2',
                        trust_remote_code=True,
                        use_safetensors=True,
                        torch_dtype=dtype,
                        low_cpu_mem_usage=True
                    )
                except:
                    model = AutoModel.from_pretrained(
                        MODEL_PATH,
                        attn_implementation='eager',
                        trust_remote_code=True,
                        use_safetensors=True,
                        torch_dtype=dtype,
                        low_cpu_mem_usage=True
                    )
                
                model = model.eval().cuda()
                self.model = model
                self.tokenizer = tokenizer
                
                self.status_var.set(f"Ready - GPU: {torch.cuda.get_device_name(0)}")
                self.process_btn.config(state=tk.NORMAL)
            except Exception as e:
                messagebox.showerror("Error", f"Failed to load model: {e}")
                self.status_var.set("Error loading model")
        
        thread = threading.Thread(target=load, daemon=True)
        thread.start()
    
    def upload_pdf(self):
        """Upload a PDF and convert pages to images"""
        file_path = filedialog.askopenfilename(
            title="Select PDF Document",
            filetypes=[
                ("PDF files", "*.pdf"),
                ("All files", "*.*")
            ]
        )
        
        if file_path:
            try:
                # Clear previous pages
                self.pdf_pages = []
                self.current_page = 0
                
                # Create temp directory for page images
                temp_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "temp", "pdf_pages")
                os.makedirs(temp_dir, exist_ok=True)
                
                # Clean old temp images
                for old_file in os.listdir(temp_dir):
                    os.remove(os.path.join(temp_dir, old_file))
                
                # Open PDF and convert each page to an image
                pdf_doc = fitz.open(file_path)
                total_pages = len(pdf_doc)
                
                for page_num in range(total_pages):
                    page = pdf_doc[page_num]
                    # Render at 200 DPI for good OCR quality
                    mat = fitz.Matrix(200 / 72, 200 / 72)
                    pix = page.get_pixmap(matrix=mat)
                    img_path = os.path.join(temp_dir, f"page_{page_num + 1}.png")
                    pix.save(img_path)
                    self.pdf_pages.append(img_path)
                
                pdf_doc.close()
                
                # Set the first page as the current image
                self.image_path = self.pdf_pages[0]
                self.current_page = 0
                
                # Display the first page
                self.display_page(0)
                
                # Update page navigation
                if total_pages > 1:
                    self.prev_btn.config(state=tk.DISABLED)
                    self.next_btn.config(state=tk.NORMAL)
                    self.page_label_var.set(f"Page 1 of {total_pages}")
                else:
                    self.prev_btn.config(state=tk.DISABLED)
                    self.next_btn.config(state=tk.DISABLED)
                    self.page_label_var.set(f"Page 1 of 1")
                
                self.status_var.set(f"Loaded: {os.path.basename(file_path)} ({total_pages} pages)")
            except Exception as e:
                messagebox.showerror("Error", f"Failed to load PDF: {e}")
    
    def display_page(self, page_index):
        """Display a specific PDF page in the preview"""
        if 0 <= page_index < len(self.pdf_pages):
            img = Image.open(self.pdf_pages[page_index])
            display_size = (440, 550)
            img.thumbnail(display_size, Image.Resampling.LANCZOS)
            photo = ImageTk.PhotoImage(img)
            self.image_label.config(image=photo, text="")
            self.image_label.image = photo
    
    def prev_page(self):
        """Navigate to the previous PDF page"""
        if self.current_page > 0:
            self.current_page -= 1
            self.display_page(self.current_page)
            total = len(self.pdf_pages)
            self.page_label_var.set(f"Page {self.current_page + 1} of {total}")
            self.next_btn.config(state=tk.NORMAL)
            if self.current_page == 0:
                self.prev_btn.config(state=tk.DISABLED)
    
    def next_page(self):
        """Navigate to the next PDF page"""
        if self.current_page < len(self.pdf_pages) - 1:
            self.current_page += 1
            self.display_page(self.current_page)
            total = len(self.pdf_pages)
            self.page_label_var.set(f"Page {self.current_page + 1} of {total}")
            self.prev_btn.config(state=tk.NORMAL)
            if self.current_page == total - 1:
                self.next_btn.config(state=tk.DISABLED)
    
    def parse_table(self, table_html):
        """
        Convert HTML table markup to clean list of rows
        Input: <table><tr><td>Cell1</td><td>Cell2</td></tr></table>
        Output: [["Cell1", "Cell2"], ...]
        """
        rows = []
        # Find all table rows
        for row_match in re.finditer(r'<tr>(.*?)</tr>', table_html, re.DOTALL):
            # Extract cells from each row
            cells = re.findall(r'<td>(.*?)</td>', row_match.group(1))
            # Clean up cell content (remove LaTeX markers, extra spaces)
            cells = [c.strip().replace('\\(', '').replace('\\)', '').replace('$ ', '$') for c in cells]
            if any(cells):  # Only add non-empty rows
                rows.append(cells)
        return rows
    
    def extract_json(self, ocr_text):
        """
        Parse OCR markdown output into readable structured JSON format
        Output matches readable_structured.json format
        """
        # Remove debug footer if present
        if '===============save results:===============' in ocr_text:
            ocr_text = ocr_text.split('===============save results:===============')[0]
        
        # Extract raw sections first
        raw_sections = []
        title = ""
        
        # Get title
        for line in ocr_text.split('\n'):
            if line.strip().startswith('## ') and not title:
                title = line.replace('#', '').strip()
                break
        
        # Find all sections - split by <|ref|>
        parts = ocr_text.split('<|ref|>')
        
        for part in parts[1:]:
            if '<|/ref|>' not in part:
                continue
                
            # Extract type
            type_end = part.find('<|/ref|>')
            sec_type = part[:type_end].strip()
            
            # Skip if no det tags
            if '<|det|>[[' not in part or ']]<|/det|>' not in part:
                continue
            
            # Extract content
            content_start = part.find(']]<|/det|>') + 11
            content = part[content_start:].strip()
            
            # Stop at next section
            if '<|ref|>' in content:
                content = content.split('<|ref|>')[0].strip()
            
            # Clean content
            content = re.sub(r'^##\s*', '', content, flags=re.MULTILINE)
            
            raw_sections.append({
                "type": sec_type,
                "content": content
            })
        
        # Convert to readable format
        result = {
            "document_title": title,
            "content": []
        }
        
        current_section = None
        
        for item in raw_sections:
            item_type = item["type"]
            content = item["content"]
            
            if item_type == "sub_title":
                # Start new section
                current_section = {
                    "heading": content,
                    "items": []
                }
                result["content"].append(current_section)
            
            elif item_type == "table":
                # Parse table and convert to financial_table format
                table_data = self.parse_table(content)
                
                if table_data and len(table_data[0]) == 2:
                    # Two-column financial table
                    table_obj = {
                        "type": "financial_table",
                        "entries": []
                    }
                    
                    for row in table_data:
                        if len(row) == 2:
                            label = row[0].strip()
                            value = row[1].strip()
                            
                            if label and value:
                                table_obj["entries"].append({
                                    "label": label,
                                    "amount": value
                                })
                            elif label:  # Header row
                                table_obj["entries"].append({
                                    "label": label,
                                    "amount": None
                                })
                    
                    if current_section:
                        current_section["items"].append(table_obj)
                    else:
                        result["content"].append(table_obj)
                else:
                    # Generic table
                    table_obj = {
                        "type": "table",
                        "data": table_data
                    }
                    if current_section:
                        current_section["items"].append(table_obj)
                    else:
                        result["content"].append(table_obj)
            
            elif item_type == "text":
                # Regular text
                if content:
                    text_obj = {"type": "text", "value": content}
                    if current_section:
                        current_section["items"].append(text_obj)
                    else:
                        result["content"].append(text_obj)
        
        return result
    
    def clean_to_markdown(self, raw_output):
        """
        Convert raw OCR output to clean Markdown/HTML format
        Removes tags and bounding boxes, keeps headings and tables
        """
        # Remove debug lines
        lines = raw_output.split('\n')
        clean_lines = []
        for line in lines:
            # Skip debug/technical lines
            if any(x in line for x in ['BASE:', 'PATCHES:', '====', 'attention', 'position_ids', 'RoPE']):
                continue
            clean_lines.append(line)
        
        text = '\n'.join(clean_lines)
        
        # Remove all <|ref|>, <|det|>, and coordinate tags
        text = re.sub(r'<\|ref\|>.*?<\|/ref\|>', '', text)
        text = re.sub(r'<\|det\|>\[\[.*?\]\]<\|/det\|>', '', text)
        
        # Clean up extra newlines
        text = re.sub(r'\n{3,}', '\n\n', text)
        
        return text.strip()
    
    def process_image(self):
        """Process all PDF pages and extract to readable JSON"""
        if not self.pdf_pages:
            messagebox.showwarning("Warning", "Please upload a PDF first")
            return
        
        if not self.model:
            messagebox.showwarning("Warning", "Model is still loading, please wait")
            return
        
        def process():
            import io
            from contextlib import redirect_stdout
            
            started = time.perf_counter()
            self.status_var.set("Processing...")
            self.process_btn.config(state=tk.DISABLED)
            
            try:
                all_raw = []
                all_clean = []
                all_json_content = []
                total_pages = len(self.pdf_pages)
                
                for page_idx, page_path in enumerate(self.pdf_pages):
                    self.status_var.set(f"Processing page {page_idx + 1} of {total_pages}...")
                    
                    # Capture stdout to get the model output
                    stdout_capture = io.StringIO()
                    
                    with torch.no_grad():
                        with redirect_stdout(stdout_capture):
                            self.model.infer(
                                self.tokenizer,
                                prompt="<image>\n<|grounding|>Convert the document to markdown.",
                                image_file=page_path,
                                output_path=str(PROJECT_DIR / "temp"),
                                base_size=1024,
                                image_size=768,
                                crop_mode=True,
                                save_results=False
                            )
                    
                    captured = stdout_capture.getvalue()
                    
                    # Collect raw output per page
                    all_raw.append(f"{'='*60}\n📄 PAGE {page_idx + 1} of {total_pages}\n{'='*60}\n{captured}")
                    
                    # Collect clean markdown per page
                    clean_text = self.clean_to_markdown(captured)
                    all_clean.append(f"## 📄 Page {page_idx + 1}\n\n{clean_text}")
                    
                    # Collect JSON per page
                    page_result = self.extract_json(captured)
                    all_json_content.append({
                        "page": page_idx + 1,
                        "document_title": page_result.get("document_title", ""),
                        "content": page_result.get("content", [])
                    })
                
                # === OUTPUT 1: RAW ===
                self.raw_output.delete(1.0, tk.END)
                self.raw_output.insert(1.0, "\n\n".join(all_raw))
                
                # === OUTPUT 2: CLEAN MARKDOWN ===
                self.clean_output.delete(1.0, tk.END)
                self.clean_output.insert(1.0, "\n\n---\n\n".join(all_clean))
                
                # === OUTPUT 3: JSON ===
                merged_json = {
                    "total_pages": total_pages,
                    "pages": all_json_content
                }
                json_str = json.dumps(merged_json, indent=2, ensure_ascii=False)
                self.json_output.delete(1.0, tk.END)
                self.json_output.insert(1.0, json_str)
                
                if self.output_dir:
                    self.output_dir.mkdir(parents=True, exist_ok=True)
                    (self.output_dir / "raw.txt").write_text("\n\n".join(all_raw), encoding="utf-8")
                    (self.output_dir / "extracted.md").write_text("\n\n---\n\n".join(all_clean), encoding="utf-8")
                    (self.output_dir / "structured.json").write_text(json_str, encoding="utf-8")
                    metadata = {"total_pages": total_pages, "processing_seconds": round(time.perf_counter() - started, 2), "gpu": torch.cuda.get_device_name(0), "model": "DeepSeek-OCR-2", "base_size": 1024, "image_size": 768, "crop_mode": True, "render_dpi": 200}
                    (self.output_dir / "run-metadata.json").write_text(json.dumps(metadata, indent=2), encoding="utf-8")
                self.notebook.select(1)
                self.status_var.set(f"Processing complete! ({total_pages} pages, {time.perf_counter() - started:.1f}s)")
                self.process_btn.config(state=tk.NORMAL)
            except Exception as e:
                messagebox.showerror("Error", f"Processing failed: {e}")
                self.status_var.set("Error")
                self.process_btn.config(state=tk.NORMAL)
        
        thread = threading.Thread(target=process, daemon=True)
        thread.start()
    
    def copy_tab(self, tab_name):
        """Copy content from specific tab to clipboard"""
        if tab_name == "raw":
            content = self.raw_output.get(1.0, tk.END).strip()
        elif tab_name == "clean":
            content = self.clean_output.get(1.0, tk.END).strip()
        elif tab_name == "json":
            content = self.json_output.get(1.0, tk.END).strip()
        else:
            return
        
        if content:
            self.root.clipboard_clear()
            self.root.clipboard_append(content)
            self.status_var.set(f"✓ Copied {tab_name.upper()} to clipboard")
    
    def save_tab(self, tab_name):
        """Save content from specific tab to file"""
        if tab_name == "raw":
            content = self.raw_output.get(1.0, tk.END).strip()
            default_ext = ".txt"
            filetypes = [("Text files", "*.txt"), ("All files", "*.*")]
        elif tab_name == "clean":
            content = self.clean_output.get(1.0, tk.END).strip()
            default_ext = ".md"
            filetypes = [("Markdown files", "*.md"), ("Text files", "*.txt"), ("All files", "*.*")]
        elif tab_name == "json":
            content = self.json_output.get(1.0, tk.END).strip()
            default_ext = ".json"
            filetypes = [("JSON files", "*.json"), ("All files", "*.*")]
        else:
            return
        
        if not content:
            messagebox.showwarning("Warning", "No content to save")
            return
        
        file_path = filedialog.asksaveasfilename(
            title=f"Save {tab_name.upper()} Output",
            defaultextension=default_ext,
            filetypes=filetypes
        )
        
        if file_path:
            try:
                with open(file_path, 'w', encoding='utf-8') as f:
                    f.write(content)
                self.status_var.set(f"✓ Saved to {os.path.basename(file_path)}")
            except Exception as e:
                messagebox.showerror("Error", f"Failed to save: {e}")
    
    def export_excel(self):
        """Export the structured JSON data to an Excel (.xlsx) file"""
        json_content = self.json_output.get(1.0, tk.END).strip()
        if not json_content:
            messagebox.showwarning("Warning", "No data to export. Process a PDF first.")
            return
        
        try:
            data = json.loads(json_content)
        except json.JSONDecodeError:
            messagebox.showerror("Error", "Invalid JSON data in output.")
            return
        
        file_path = filedialog.asksaveasfilename(
            title="Export to Excel",
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")]
        )
        
        if not file_path:
            return
        
        try:
            wb = Workbook()
            # Remove default sheet
            wb.remove(wb.active)
            
            # Style definitions
            header_font = Font(bold=True, size=12, color="FFFFFF")
            header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            section_font = Font(bold=True, size=11, color="1F4E79")
            section_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
            thin_border = Border(
                left=Side(style='thin'), right=Side(style='thin'),
                top=Side(style='thin'), bottom=Side(style='thin')
            )
            
            pages = data.get("pages", [])
            if not pages:
                # Single-page fallback (old format without "pages" key)
                pages = [{"page": 1, "document_title": data.get("document_title", ""), "content": data.get("content", [])}]
            
            for page_data in pages:
                page_num = page_data.get("page", 1)
                doc_title = page_data.get("document_title", f"Page {page_num}")
                sheet_name = f"Page {page_num}"
                
                # Excel sheet names max 31 chars
                ws = wb.create_sheet(title=sheet_name[:31])
                row = 1
                
                # Write document title
                ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
                cell = ws.cell(row=row, column=1, value=doc_title)
                cell.font = Font(bold=True, size=14)
                cell.alignment = Alignment(horizontal='center')
                row += 2
                
                content = page_data.get("content", [])
                
                for section in content:
                    # Section with heading
                    heading = section.get("heading", "")
                    if heading:
                        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
                        cell = ws.cell(row=row, column=1, value=heading)
                        cell.font = section_font
                        cell.fill = section_fill
                        cell.border = thin_border
                        row += 1
                    
                    items = section.get("items", [])
                    # If it's a direct item (no heading/items structure)
                    if not items and not heading:
                        items = [section]
                    
                    for item in items:
                        item_type = item.get("type", "")
                        
                        if item_type in ("financial_table", "table"):
                            entries = item.get("entries", [])
                            table_data = item.get("data", [])
                            
                            if entries:
                                # Financial table: Label | Amount
                                # Header row
                                for col_idx, col_name in enumerate(["Label", "Amount"], 1):
                                    cell = ws.cell(row=row, column=col_idx, value=col_name)
                                    cell.font = header_font
                                    cell.fill = header_fill
                                    cell.border = thin_border
                                    cell.alignment = Alignment(horizontal='center')
                                row += 1
                                
                                for entry in entries:
                                    label = entry.get("label", "")
                                    amount = entry.get("amount", "")
                                    ws.cell(row=row, column=1, value=label).border = thin_border
                                    ws.cell(row=row, column=2, value=amount if amount else "").border = thin_border
                                    row += 1
                            
                            elif table_data:
                                # Generic table: write as-is
                                for row_data in table_data:
                                    for col_idx, cell_val in enumerate(row_data, 1):
                                        cell = ws.cell(row=row, column=col_idx, value=cell_val)
                                        cell.border = thin_border
                                    row += 1
                        
                        elif item_type == "text":
                            value = item.get("value", "")
                            if value:
                                ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
                                ws.cell(row=row, column=1, value=value).alignment = Alignment(wrap_text=True)
                                row += 1
                        
                        row += 1  # Space between items
                
                # Auto-fit column widths (approximate)
                from openpyxl.utils import get_column_letter
                for col_idx in range(1, ws.max_column + 1):
                    max_length = 0
                    for row_idx in range(1, ws.max_row + 1):
                        cell = ws.cell(row=row_idx, column=col_idx)
                        try:
                            if cell.value and not isinstance(cell, type(None)):
                                max_length = max(max_length, len(str(cell.value)))
                        except:
                            pass
                    col_letter = get_column_letter(col_idx)
                    ws.column_dimensions[col_letter].width = min(max_length + 4, 50)
            
            wb.save(file_path)
            self.status_var.set(f"✓ Excel exported to {os.path.basename(file_path)}")
            messagebox.showinfo("Success", f"Excel file saved to:\n{file_path}")
        
        except Exception as e:
            messagebox.showerror("Error", f"Failed to export Excel: {e}")
    
    def clear_all(self):
        """Clear all inputs and outputs"""
        self.image_label.config(image='', text="No document loaded")
        self.image_label.image = None
        self.raw_output.delete(1.0, tk.END)
        self.clean_output.delete(1.0, tk.END)
        self.json_output.delete(1.0, tk.END)
        self.image_path = None
        self.pdf_pages = []
        self.current_page = 0
        self.page_label_var.set("")
        self.prev_btn.config(state=tk.DISABLED)
        self.next_btn.config(state=tk.DISABLED)
        self.status_var.set("Ready")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Local PDF OCR with DeepSeek-OCR-2")
    parser.add_argument("--output-dir", help="Automatically save raw, Markdown, JSON and run metadata after processing")
    args = parser.parse_args()
    root = tk.Tk()
    app = OCRApp(root, output_dir=args.output_dir)
    root.mainloop()
